from rest_framework.decorators import api_view
from rest_framework.response import Response
from .models import *
from .serializers import *
from django.contrib.auth import login as django_login
from django.contrib.auth.models import User
from django.contrib.auth.hashers import check_password
from django.shortcuts import get_object_or_404
from rest_framework.decorators import permission_classes
from openpyxl import load_workbook
from django.conf import settings
from django.db import transaction
import os

# Create your views here.

# HELPER: Check Admin user_id --------------------------------
def get_valid_admin(request):
    user_id = request.session.get('user_master_id')
    if not user_id:
        return None, Response({"error": "Please login first"}, status=401)
    admin_user = UserMaster.objects.filter(id=user_id, account_type='Admin', active_status=True).first()
    if not admin_user:
        return None, Response({"error": "Not authorized as Admin"}, status=403)
    return admin_user, None

# CUSTOMER LOGIN ------------------------------------------------
@api_view(['POST'])
def customer_login(request):
    username = request.data.get('username')
    password = request.data.get('password')

    if not username or not password:
        return Response({"error": "Username and Password are required"}, status=400)

    try:
        user = UserMaster.objects.get(username=username)
    except UserMaster.DoesNotExist:
        return Response({"error": "Invalid Username or Password"}, status=401)

    if user.account_type != "Standard":
        return Response({"error": "Only Customer can login"}, status=403)

    if not user.active_status:
        return Response({"error": "User account is deactivated"}, status=403)

    if not check_password(password, user.password):
        return Response({"error": "Invalid Username or Password"}, status=401)

    # AUTOMATIC SESSION CREATE ---------------------
    django_user, created = User.objects.get_or_create(username=user.username)
    django_login(request, django_user)

    request.session['user_master_id'] = user.id  # yahi automatically future requests me milega

    return Response({
        "message": "Customer Login Successful",
        "customer": {
            "id": user.id,
            "username": user.username,
            "name": user.name,
            "account_type": user.account_type
        }
    })

# ADMIN LOGIN ------------------------------------------------
@api_view(['POST'])
def admin_login(request):

    username = request.data.get('username')
    password = request.data.get('password')

    if not username or not password:
        return Response({
            "error": "Username and Password are required"
        }, status=400)

    try:
        user = UserMaster.objects.get(username=username)
    except UserMaster.DoesNotExist:
        return Response({
            "error": "Invalid Username or Password"
        }, status=401)

    if user.account_type != "Admin":
        return Response({
            "error": "Only Admin can login"
        }, status=403)

    if not user.active_status:
        return Response({
            "error": "User account is deactivated"
        }, status=403)

    if not check_password(password, user.password):
        return Response({
            "error": "Invalid Username or Password"
        }, status=401)

    django_user, created = User.objects.get_or_create(username=user.username)
    django_login(request, django_user)
    request.session['user_master_id'] = user.id

    return Response({
        "message": "Admin Login Successful",
        "admin": {
            "id": user.id,
            "username": user.username,
            "name": user.name,
            "account_type": user.account_type
        }
    })

# GET ALL USERS -----------------------------------------
@api_view(['GET'])
def get_users(request):
    admin_user, error = get_valid_admin(request)
    if error:
        return error
    users = UserMaster.objects.all().values('id', 'username', 'name', 'account_type', 'active_status')
    return Response(users)

# ADD USER ----------------------------------------------
@api_view(['POST'])
def add_user(request):
    username = request.data.get('username')
    password = request.data.get('password')
    name = request.data.get('name')
    account_type = request.data.get('account_type')

    if not username or not password or not name:
        return Response({"error": "username, password and name are required"}, status=400)

    if UserMaster.objects.filter(username=username).exists():
        return Response({"error": "Username already exists"})

    if account_type == 'Admin':
        admin_exists = UserMaster.objects.filter(account_type='Admin').exists()
        if admin_exists:
            admin_user, error = get_valid_admin(request)
            if error:
                return error

    user = UserMaster.objects.create(
        username=username,
        password=password,
        name=name,
        account_type=account_type)

    return Response({
        "message": "User Added Successfully",
        "username": user.username})

# GET SINGLE USER -----------------------------------------
@api_view(['GET'])
def get_single_user(request, id):
    admin_user, error = get_valid_admin(request)
    if error:
        return error
    user = get_object_or_404(UserMaster, id=id)
    serializer = UserSerializer(user)
    return Response(serializer.data)

# UPDATE USER -----------------------------------------
@api_view(['PUT'])
def update_user(request, id):
    admin_user, error = get_valid_admin(request)
    if error:
        return error
    user_master = get_object_or_404(UserMaster, id=id)
    username = request.data.get("username", user_master.username)
    name = request.data.get("name", user_master.name)
    account_type = request.data.get("account_type", user_master.account_type)
    user_master.username = username
    user_master.name = name
    user_master.account_type = account_type
    if request.data.get("password"):
        user_master.password = request.data.get("password")
    user_master.save()
    return Response({"message": "User Updated Successfully"})

# UPDATE USER STATUS -----------------------------------------
@api_view(['PUT'])
def update_user_status(request, id):
    admin_user, error = get_valid_admin(request)
    if error:
        return error
    user_master = get_object_or_404(UserMaster, id=id)
    status = request.data.get("active_status")
    if status is None:
        return Response({"error": "active_status field is required"})
    user_master.active_status = status
    user_master.save()
    return Response({
        "message": "User Status Updated Successfully",
        "active_status": status
    })

# GET ALL PRODUCTS-----------------------------------------------
@api_view(['GET'])
def get_products(request):
    products = ProductMaster.objects.all()
    serializer = ProductSerializer(products, many=True)
    return Response(serializer.data)

# ADD PRODUCT------------------------------------------------
@api_view(['POST'])
def add_product(request):
   admin_user, error = get_valid_admin(request)
   if error:
       return error
   serializer = ProductSerializer(data=request.data)
   if serializer.is_valid():
        serializer.save()
        return Response({'message': 'Product Added Successfully'})
   return Response(serializer.errors)

# GET SINGLE PRODUCT-------------------------------------------
@api_view(['GET'])
def get_single_product(request, id):
    product = get_object_or_404(ProductMaster,id=id,is_active=True)
    serializer = ProductSerializer(product)
    return Response(serializer.data)

# UPDATE PRODUCT-------------------------------------------------
@api_view(['PUT'])
def update_product(request, id):
    admin_user, error = get_valid_admin(request)
    if error:
       return error
    product = get_object_or_404(ProductMaster,id=id)
    serializer = ProductSerializer(product, data=request.data,partial=True)
    if serializer.is_valid():
        serializer.save()
        return Response({'message': 'Product Updated Successfully'})
    return Response(serializer.errors)

#PRODUCT STATUS------------------------------------------------
@api_view(['PUT'])
def update_product_status(request, id):
    admin_user, error = get_valid_admin(request)
    if error:
       return error
    product = get_object_or_404(ProductMaster,id=id)
    is_active = request.data.get('is_active')
    if is_active is None:return Response({"error":"is_active field is required"},status=400)
    product.is_active = is_active
    product.save()
    return Response({
        'message':
        'Product status updated successfully',
        'is_active':
        product.is_active
    })

# ADD PAYMENT MODE --------------------------------
@api_view(['POST'])
def add_payment_mode(request):
    admin_user, error = get_valid_admin(request)
    if error:
        return error
    raw_name = request.data.get("name")
    if not raw_name:
        return Response({"error": "name is required"}, status=400)
    name = " ".join(raw_name.split()).strip()
    if PaymentModeMaster.objects.filter(name__iexact=name).exists():
       return Response({"error": "Payment Mode Already Exists"})
    payment_mode = PaymentModeMaster.objects.create(name=name)
    return Response({
        "message": "Payment Mode Added Successfully",
        "name": payment_mode.name})

# GET ALL PAYMENT MODES --------------------------------
@api_view(['GET'])
def get_payment_modes(request):
    payment_modes = PaymentModeMaster.objects.all().values()
    return Response(payment_modes)

# GET SINGLE PAYMENT MODE -----------------------------------------
@api_view(['GET'])
def get_single_payment_mode(request, id):
    payment_mode = get_object_or_404(PaymentModeMaster, payment_mode_id=id)
    serializer = PaymentModeSerializer(payment_mode)
    return Response(serializer.data)

# UPDATE PAYMENT MODE -----------------------------------------
@api_view(['PUT'])
def update_payment_mode(request, id):
    admin_user, error = get_valid_admin(request)
    if error:
        return error
    payment_mode = get_object_or_404(
        PaymentModeMaster,
        payment_mode_id=id)
    serializer = PaymentModeSerializer(
        payment_mode,
        data=request.data,
        partial=True)
    if serializer.is_valid():
        serializer.save()
        return Response({
            "message": "Payment Mode Updated Successfully"})
    return Response(serializer.errors)

# UPDATE PAYMENT MODE STATUS -----------------------------------------
@api_view(['PUT'])
def update_payment_mode_status(request, id):
    admin_user, error = get_valid_admin(request)
    if error:
        return error
    payment_mode = get_object_or_404(
        PaymentModeMaster,
        payment_mode_id=id)
    status = request.data.get("active_status")
    if status is None:
        return Response({
            "error": "active_status field is required"})
    payment_mode.active_status = status
    payment_mode.save()
    return Response({
        "message": "Payment Mode Status Updated Successfully",
        "active_status": payment_mode.active_status})

# IMPORT PRODUCTS FROM EXCEL ------------------------------------------
@transaction.atomic
@api_view(['POST'])
def import_products(request):
    admin_user, error = get_valid_admin(request)
    if error:
        return error
    if 'file' not in request.FILES:
        return Response({"error": "Excel file is required"})
    excel_file = request.FILES['file']
    try:
        workbook = load_workbook(excel_file)
        sheet = workbook.active
        imported = 0
        skipped = 0
        errors = []
        # Skip Header Row
        for row in sheet.iter_rows(min_row=2, values_only=True):
            name = " ".join(str(row[0]).split())
            sale_price = row[1]
            stock = row[2]
            photo = row[3]
            is_active = row[4]
            # Required Fields
            if not name or sale_price is None or stock is None:
                skipped += 1
                errors.append(f"{name} -> Missing Required Data")
                continue
            # Duplicate Product
            if ProductMaster.objects.filter(name__iexact=name).exists():
               skipped += 1
               errors.append(f"{name} -> Already Exists")
               continue
            product = ProductMaster(
                name=name,
                sale_price=sale_price,
                stock=stock,
                is_active=True if is_active in [True, "True", 1, "1", None] else False)
            # Optional Photo
            if photo:
                image_path = os.path.join(settings.MEDIA_ROOT,"products",str(photo))
                if os.path.exists(image_path):
                    product.photo = f"products/{photo}"
            product.save()
            imported += 1
        return Response({
            "message": "Products Imported Successfully",
            "Imported Products": imported,
            "Skipped Products": skipped,
            "Errors": errors})
    except Exception as e:
        return Response({"error": str(e)})